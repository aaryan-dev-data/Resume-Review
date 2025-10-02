import streamlit as st
from transformers import pipeline
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
import re

# ----------------- Helper Functions -----------------

def extract_text(file):
    """Extract text from PDF or DOCX."""
    if file.type == "application/pdf":
        pdf = PdfReader(file)
        text = ""
        for page in pdf.pages:
            text += page.extract_text() + "\n"
        return text
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(file)
        return "\n".join([p.text for p in doc.paragraphs])
    else:
        return ""

def split_sections(text):
    """Split text into common resume sections."""
    sections = {}
    current_sec = "SUMMARY"
    sections[current_sec] = ""
    lines = text.split("\n")
    for line in lines:
        header = line.strip().upper()
        if header in ["SUMMARY", "EXPERIENCE", "SKILLS", "PROJECTS", "EDUCATION"]:
            current_sec = header
            sections[current_sec] = ""
        else:
            sections[current_sec] += line + "\n"
    return sections

def chunk_text(text, max_tokens=300):
    """Split large text into chunks for the model."""
    words = text.split()
    chunks = []
    for i in range(0, len(words), max_tokens):
        chunks.append(" ".join(words[i:i+max_tokens]))
    return chunks

def rewrite_section(section_name, section_text, analyzer):
    """Rewrite a section professionally and ATS-friendly."""
    if not section_text.strip():
        return f"{section_name} section is missing or empty."
    chunks = chunk_text(section_text, max_tokens=300)
    feedback = ""
    for chunk in chunks:
        prompt = f"""
Rewrite this {section_name} section for a professional, ATS-friendly resume.
- Use impact-oriented language and action verbs
- Add metrics where possible
- Group skills logically
- Make it concise and professional
Section:
{chunk}
"""
        try:
            response = analyzer(prompt, max_length=350, do_sample=True, temperature=0.7, top_p=0.95)[0]['generated_text']
            feedback += response.strip() + "\n"
        except Exception as e:
            feedback += f"Error analyzing chunk: {str(e)}\n"
    return feedback.strip()

def generate_feedback(text, analyzer):
    """Generate rewritten feedback for all sections."""
    sections = split_sections(text)
    feedback = {}
    for sec, content in sections.items():
        feedback[sec] = rewrite_section(sec, content, analyzer)
    return feedback

def infer_role_keywords(text, analyzer):
    """Infer likely role and suggest ATS keywords."""
    prompt_role = f"Analyze this resume and infer the most probable job role. Reply only with role name: {text[:1000]}"
    role = analyzer(prompt_role, max_length=20, do_sample=False)[0]['generated_text'].strip()
    
    prompt_keywords = f"List 10–15 important keywords for a {role} role. Reply with comma-separated keywords only."
    response = analyzer(prompt_keywords, max_length=60, do_sample=False)[0]['generated_text']
    keywords = [kw.strip() for kw in response.split(",") if kw.strip()]
    return role, keywords

def check_ats_keywords(text, keywords):
    present = [kw for kw in keywords if re.search(rf"\b{kw}\b", text, re.I)]
    missing = [kw for kw in keywords if kw not in present]
    return present, missing

def create_pdf(feedback, role=None, present=None, missing=None, filename="Optimized_Resume.pdf"):
    """Generate PDF with rewritten sections and ATS feedback."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, "Optimized Resume Feedback", ln=True, align="C")
    pdf.ln(5)
    
    pdf.set_font("Arial", '', 12)
    for sec, fb in feedback.items():
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 8, sec, ln=True)
        pdf.set_font("Arial", '', 12)
        for line in fb.split("\n"):
            if line.strip():
                pdf.multi_cell(0, 6, "- " + line.strip())
        pdf.ln(3)
    
    if role:
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 8, f"Inferred Role: {role}", ln=True)
        pdf.set_font("Arial", '', 12)
        pdf.multi_cell(0, 6, f"✅ Present Keywords: {', '.join(present) if present else 'None'}")
        pdf.multi_cell(0, 6, f"❌ Missing Keywords: {', '.join(missing) if missing else 'None'}")
    
    pdf.output(filename)
    return filename

# ----------------- Streamlit UI -----------------

st.title("AI Resume Optimizer – ATS-friendly Rewrite + PDF Export")

uploaded_file = st.file_uploader("Upload your Resume (PDF or DOCX)", type=["pdf", "docx"])

if uploaded_file:
    st.info("Extracting text from resume...")
    resume_text = extract_text(uploaded_file)
    st.success("Text extraction completed!")
    
    st.info("Analyzing and rewriting resume sections... This may take 20–40 seconds on first run.")
    
    analyzer = pipeline("text2text-generation", model="google/flan-t5-base")
    
    # Rewriting all sections
    feedback = generate_feedback(resume_text, analyzer)
    
    st.subheader("Section-wise Feedback")
    for sec, fb in feedback.items():
        st.markdown(f"**{sec}**")
        st.write(fb)
    
    # ATS Feedback
    st.subheader("ATS Feedback")
    role, keywords = infer_role_keywords(resume_text, analyzer)
    present, missing = check_ats_keywords(resume_text, keywords)
    st.write(f"Inferred Role: **{role}**")
    st.write(f"✅ Present Keywords: {', '.join(present) if present else 'None'}")
    st.write(f"❌ Missing Keywords: {', '.join(missing) if missing else 'None'}")
    
    # PDF Export
    st.subheader("Download Optimized Resume PDF")
    if st.button("Generate PDF"):
        pdf_file = create_pdf(feedback, role, present, missing)
        with open(pdf_file, "rb") as f:
            st.download_button(
                label="Download PDF",
                data=f,
                file_name="Optimized_Resume.pdf",
                mime="application/pdf"
            )
        st.success("PDF generated successfully!")
